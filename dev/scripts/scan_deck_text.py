# -*- coding: utf-8 -*-
"""① 打印 6~24 页的文字（看用户改了什么）
   ② 自动扫一遍客观问题：** 残留、孤立星号、连续空格、全半角括号混用、
      超长行（可能溢出）、重复标点、中英标点混用、占位符残留
"""
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

ROOT = r"E:\university\order\python\pet\AIpet-Murasame-main"
DECK = os.path.join(ROOT, "dev", "ppt", "out", "Deck.pptx")

prs = Presentation(DECK)
slides = list(prs.slides)

print("=== ① 第 6~24 页文字 ===")
for i, slide in enumerate(slides, 1):
    if i < 6:
        continue
    lines = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    lines.append(t)
    if lines:
        print("   ── 第 %d 页 ──" % i)
        for t in lines:
            print("      %s" % t[:112])

print()
print("=== ② 客观问题扫描 ===")
ISSUES = [
    (r"\*\*", "** 残留（Markdown 没被引擎解析）"),
    (r"(?<!\*)\*(?!\*)", "孤立星号 *"),
    (r"  +", "连续多个空格"),
    (r"[（(][^）)]*$", "括号没闭合"),
    (r"[，。；：]{2,}", "重复中文标点"),
    (r"[,.;:]{2,}", "重复英文标点"),
    (r"[\u4e00-\u9fff],", "中文后用英文逗号"),
    (r"[\u4e00-\u9fff]\.(?![0-9a-zA-Z])", "中文后用英文句点"),
    (r"TODO|TBD|待补|xxx|XXX|占位|placeholder", "占位符残留"),
    (r"\\n", "字面 \\n"),
    (r"&amp;|&lt;|&gt;", "HTML 实体未转义"),
]
found_any = False
for i, slide in enumerate(slides, 1):
    for sh in slide.shapes:
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        txt = sh.text_frame.text
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            if not t.strip():
                continue
            for pat, desc in ISSUES:
                if re.search(pat, t):
                    found_any = True
                    m = re.search(pat, t)
                    print("   第 %2d 页 [%s] …%s…" % (i, desc, t[max(0, m.start() - 18):m.end() + 18].strip()))
if not found_any:
    print("   ✓ 没扫出上述问题")

print()
print("=== ③ 行长统计（看有没有可能溢出的超长段落）===")
rows = []
for i, slide in enumerate(slides, 1):
    for sh in slide.shapes:
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                rows.append((len(t), i, t))
rows.sort(reverse=True)
for n, i, t in rows[:14]:
    print("   %3d 字  第 %2d 页  %s" % (n, i, t[:96]))

print()
print("=== ④ 全角/半角括号与引号统计 ===")
alltext = []
for slide in slides:
    for sh in slide.shapes:
        if sh.has_text_frame:
            alltext.append(sh.text_frame.text)
s = "\n".join(alltext)
for name, ch in (("（全角左括号", "（"), ("(半角左括号", "("), ("“左双引号", "“"),
                 ("\"直双引号", '"'), ("「左角括号", "「"), ("▪前缀", "▪"), ("·间隔号", "·")):
    print("   %-12s %d 次" % (name, s.count(ch)))
