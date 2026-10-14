# -*- coding: utf-8 -*-
"""① 补 PPT 第 21 页的"怎么才能真看到"与更强的真机措辞
   ② 找"规划文档"（未来规划/路线图）在哪，并打印其结构
"""
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

ROOT = r"E:\university\order\python\pet\AIpet-Murasame-main"
DECK = os.path.join(ROOT, "dev", "ppt", "out", "Deck.pptx")

print("=== ① 补第 21 页 ===")
prs = Presentation(DECK)
slides = list(prs.slides)
before = [[sh.left, sh.top, sh.width, sh.height] for s in slides for sh in s.shapes
          if sh.shape_type == 13]
slide = slides[20]
NOTE = ("；**真机还没用真手验过** —— 要用真手须先关掉界面上的「试手」（那是假手），"
        "并先把 .vrm 模型 adb push 到 files/vrm/model.vrm")
n = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        t = "".join(r.text for r in para.runs)
        newt = t
        if "代码已通：手部 21 关键点检测" in t and "试手" not in t:
            newt = t.rstrip() + NOTE
        if "未验证：真机" in t:
            newt = ("▪  未验证：真机  编译 / 21 个 JVM 单测 / JS 语法都过了，"
                    "但**真机从未用真手验过** ——「手伸过去她眼睛是否跟着转」「捏一下她是否躲」都还是未知；"
                    "真机只跑过假手的坐标链路（rotation=90，40 秒命中 4 次）")
        if newt != t:
            if len(para.runs) == 1:
                para.runs[0].text = newt
            else:
                para.runs[0].text = newt
                for r in para.runs[1:]:
                    r.text = ""
            print("   ✓ " + newt[:110])
            n += 1
if n:
    prs.save(DECK)
prs2 = Presentation(DECK)
after = [[sh.left, sh.top, sh.width, sh.height] for s in prs2.slides for sh in s.shapes
         if sh.shape_type == 13]
print("   改了 %d 段；图片 %d→%d %s" % (n, len(before), len(after),
                                       "✓✓ 完全一致" if before == after else "✗ 变了"))

print()
print("=== ② 找规划文档 ===")
PAT = re.compile(r"规划|路线|roadmap|未来|下一步", re.I)
cands = []
for sub in ("dev/docs", "docs", "."):
    d = os.path.join(ROOT, sub) if sub != "." else ROOT
    if not os.path.isdir(d):
        continue
    for f in sorted(os.listdir(d)):
        if not f.endswith(".md"):
            continue
        p = os.path.join(d, f)
        if PAT.search(f):
            cands.append((os.path.getsize(p), os.path.join(sub, f).replace("\\", "/")))
for sz, rel in sorted(cands, reverse=True)[:14]:
    print("   %8.1f KB  %s" % (sz / 1024, rel))

print()
print("=== ②-b 与未来规划最相关的那份，打印其标题结构 ===")
target = None
for sz, rel in sorted(cands, reverse=True):
    if "PPT" in rel or "规划" in rel:
        target = rel
        break
if target is None and cands:
    target = sorted(cands, reverse=True)[0][1]
print("   选中：%s" % target)
if target:
    p = os.path.join(ROOT, target.replace("/", os.sep))
    for i, l in enumerate(io.open(p, encoding="utf-8", errors="replace").read().split("\n"), 1):
        if l.startswith("#") or re.match(r"^\s*[-*]\s", l):
            print("   %4d| %s" % (i, l.rstrip()[:104]))
