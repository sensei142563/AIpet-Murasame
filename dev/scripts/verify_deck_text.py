# -*- coding: utf-8 -*-
"""只改文字，绝不碰图片。步骤：
  ① 逐页核对"页码"文字与页序是否一致（找出 20/21 互换）
  ② 定位那 1 处半角左括号 (
  ③ 核实第 15 页 affinity_floor 这个断言在代码里是否真实存在
  ④ 记录修改前后所有图片的 (左,上,宽,高) —— 用于事后断言"一个都没动"
"""
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

ROOT = r"E:\university\order\python\pet\AIpet-Murasame-main"
DECK = os.path.join(ROOT, "dev", "ppt", "out", "Deck.pptx")

prs = Presentation(DECK)
slides = list(prs.slides)

print("=== ① 页码核对 ===")
print("   页序 | 该页里的纯数字段落")
for i, slide in enumerate(slides, 1):
    nums = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if re.fullmatch(r"\d{1,3}", t):
                w = (sh.width or 0) / 914400.0
                h = (sh.height or 0) / 914400.0
                nums.append((t, round(w, 2), round(h, 2)))
    flag = ""
    for t, w, h in nums:
        if t == str(i):
            flag = "✓"
            break
    else:
        if nums:
            flag = "✗ 与页序不符"
    print("   %3d  | %s  %s" % (i, nums if nums else "（无数字）", flag))

print()
print("=== ② 半角左括号定位 ===")
for i, slide in enumerate(slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        full = sh.text_frame.text
        for m in re.finditer(r"\(", full):
            print("   第 %d 页：…%s…" % (i, full[max(0, m.start() - 30):m.start() + 30].replace("\n", " ")))

print()
print("=== ③ 核实 affinity_floor 是否真的存在 ===")
hits = []
for dp, dn, fn in os.walk(ROOT):
    if any(s in dp for s in (".git", "runtime", "GPT-SoVITS", "F5-TTS_Models",
                             "\\models", "dev\\tmp", ".venv", "__pycache__", "NapCat")):
        continue
    for f in fn:
        if not f.endswith(".py"):
            continue
        p = os.path.join(dp, f)
        try:
            t = io.open(p, encoding="utf-8", errors="replace").read()
        except Exception:
            continue
        if "affinity_floor" in t:
            rel = os.path.relpath(p, ROOT)
            for j, l in enumerate(t.split("\n"), 1):
                if "affinity_floor" in l:
                    hits.append((rel, j, l.strip()))
print("   命中 %d 处：" % len(hits))
for rel, j, l in hits[:18]:
    print("      %-38s %5d| %s" % (rel[:38], j, l[:100]))
if not hits:
    print("      ✗ 代码里**没有** affinity_floor —— 第 15 页那句断言可能是错的")
else:
    # 看有没有实现（不只是注释/文档）
    impl = [h for h in hits if h[2].startswith(("def ", "affinity_floor", "return", "if ",
                                                "floor", "raise", "print"))
            or "=" in h[2]]
    print("      ⇒ 其中像实现或赋值的 %d 处" % len(impl))

print()
print("=== ④ 图片快照（改字前后必须完全一致）===")
import json  # noqa: E402
snap = []
for i, slide in enumerate(slides, 1):
    for sh in slide.shapes:
        if sh.shape_type == 13:
            snap.append([i, sh.left, sh.top, sh.width, sh.height, sh.name])
sp = os.path.join(ROOT, "dev", "ppt", "out", "_img_snapshot_before.json")
io.open(sp, "w", encoding="utf-8", newline="\n").write(json.dumps(snap, ensure_ascii=False))
print("   %d 张图片的位置/尺寸已存 %s" % (len(snap), os.path.basename(sp)))
