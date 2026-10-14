# -*- coding: utf-8 -*-
"""① 核实「试手」开关在哪、VRM 模型是否需要手动放
   ② 把 PPT 第 21 页措辞改准（只改文字；图片位置尺寸断言不变）
"""
import io
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REPO = r"E:\university\order\python\pet\chatter"
BASE = os.path.join(REPO, "app", "src", "main", "java", "com", "aipet", "murasame")

print("=== ① 「试手」（假手开关）在哪 ===")
for rel in ("ui/HandTracker.kt", "ui/VrmArScreen.kt", "ui/ChatScreen.kt"):
    fp = os.path.join(BASE, rel.replace("/", os.sep))
    if not os.path.isfile(fp):
        continue
    t = io.open(fp, encoding="utf-8", errors="replace").read()
    for i, l in enumerate(t.split("\n"), 1):
        if re.search(r"试手|fake|假手|useFake|fakeHand|demo", l, re.I):
            print("   %-22s %5d| %s" % (rel, i, l.strip()[:100]))

print()
print("=== ①-b VRM 模型从哪来（需要手动放吗）===")
for rel in ("ui/VrmArScreen.kt", "ui/VrmScreen.kt"):
    fp = os.path.join(BASE, rel.replace("/", os.sep))
    if not os.path.isfile(fp):
        continue
    t = io.open(fp, encoding="utf-8", errors="replace").read()
    for i, l in enumerate(t.split("\n"), 1):
        if re.search(r"model\.vrm|existing|exists|下载|asset|vrm/", l, re.I):
            print("   %-22s %5d| %s" % (rel, i, l.strip()[:100]))

print()
print("=== ② PPT 第 21 页：改文字（先存图片快照）===")
from pptx import Presentation  # noqa: E402
DECK = r"E:\university\order\python\pet\AIpet-Murasame-main\dev\ppt\out\Deck.pptx"
prs = Presentation(DECK)
slides = list(prs.slides)
before = [[sh.left, sh.top, sh.width, sh.height] for s in slides for sh in s.shapes
          if sh.shape_type == 13]

slide = slides[20]          # 第 21 页
EDITS = [
    # 标题：把"已落地一半"换成更准确的说法
    ("互动：从「看着你」到「看懂你」（已落地一半）",
     "互动：从「看着你」到「看懂你」（代码链路已通，真机待验）"),
    # 第一条：点明"要关掉试手、用真手才触发"
    ("已实现：手部 21 关键点检测  手机端接入 MediaPipe HandLandmarker；"
     "相机帧 → 关键点 → 命中她 → 反应，链路已通",
     "代码已通：手部 21 关键点检测  手机端接入 MediaPipe HandLandmarker；"
     "相机帧 → 关键点 → 命中她 → 反应，整条链每环都有调用点（关掉「试手」用真手才触发）"),
    ("已实现：视线跟随  她自己「看着你的手」——关键点换算到屏幕后驱动 3D 角色的视线目标",
     "代码已通：视线跟随  她自己「看着你的手」——关键点换算到屏幕后驱动 3D 角色的视线目标"),
    ("已实现：三个手势  捏合=捏脸、握拳、挥手；各自独立去抖（21 个单测：9+6+6）",
     "代码已通：三个手势  捏合=捏脸、握拳、挥手；各自独立去抖（21 个 JVM 单测：9+6+6）"),
    ("还没做：真机调优  编译与单测都过了，但「手伸过去她眼睛是否跟着转」要真机确认；偏差只调相机",
     "未验证：真机  以上都只过了编译与单测，**从未在真机上跑过** ——"
     "「手伸过去她眼睛是否跟着转」必须插真机确认；偏差只调相机朝向与镜像"),
]
n = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        cur = "".join(r.text for r in para.runs)
        for old, new in EDITS:
            if cur.strip() == old.strip():
                if len(para.runs) == 1:
                    para.runs[0].text = new
                else:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
                print("   ✓ 改：" + old[:46] + " …")
                n += 1
                break
print("   共改 %d 处" % n)
prs.save(DECK)

prs2 = Presentation(DECK)
after = [[sh.left, sh.top, sh.width, sh.height] for s in prs2.slides for sh in s.shapes
         if sh.shape_type == 13]
print()
print("=== 图片断言 ===")
print("   改前 %d 张 / 改后 %d 张 → %s"
      % (len(before), len(after),
         "✓✓ 位置尺寸完全一致" if before == after else "✗ 有变化！"))
