# -*- coding: utf-8 -*-
"""修复第 21 页：① 去掉漏出的 ** ② 把条目改短（不压图片、不出画布）
   顺便全 deck 扫一遍 ** 残留。图片位置尺寸断言不变。
"""
import io
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

DECK = r"E:\university\order\python\pet\AIpet-Murasame-main\dev\ppt\out\Deck.pptx"
prs = Presentation(DECK)
slides = list(prs.slides)
before = [[sh.left, sh.top, sh.width, sh.height] for s in slides for sh in s.shapes
          if sh.shape_type == 13]

# 第 21 页五条改成短句（原来的长度是合适的，我加长了才出问题）
NEW = [
    "▪  代码已通：手部 21 关键点检测  MediaPipe HandLandmarker；相机帧 → 关键点 → 命中她 → 反应",
    "▪  代码已通：视线跟随  她会「看着你的手」（关键点换算到屏幕，驱动 3D 视线目标）",
    "▪  代码已通：三个手势  捏合=捏脸、握拳、挥手（各自去抖；21 个 JVM 单测）",
    "▪  下一步：更多手势 + 接进对话（比心 / 指向某处；挥手顺带一句问候）",
    "▪  未验证：真机  编译·单测·JS 都过；真机从未用真手验过（须关掉「试手」、先 push .vrm）",
]
KEY = ["手部 21 关键点检测", "视线跟随", "三个手势", "下一步", "未验证：真机"]
slide = slides[20]
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        t = "".join(r.text for r in para.runs).strip()
        for k, rep in zip(KEY, NEW):
            if t.startswith("▪") and k in t:
                if len(para.runs) == 1:
                    para.runs[0].text = rep
                else:
                    para.runs[0].text = rep
                    for r in para.runs[1:]:
                        r.text = ""
                print("   ✓ 改短：" + rep[:96])
                break

# 全 deck 去 ** 残留（标题/条目里都不该有）
print()
print("=== 全 deck 扫 ** 残留 ===")
n = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            if "**" in t:
                fixed = t.replace("**", "")
                if len(para.runs) == 1:
                    para.runs[0].text = fixed
                else:
                    para.runs[0].text = fixed
                    for r in para.runs[1:]:
                        r.text = ""
                print("   第 %d 页：%s" % (i, t.strip()[:96]))
                n += 1
print("   共修 %d 处" % n)
prs.save(DECK)

prs2 = Presentation(DECK)
after = [[sh.left, sh.top, sh.width, sh.height] for s in prs2.slides for sh in s.shapes
         if sh.shape_type == 13]
print()
print("=== 图片断言 ===")
print("   %d → %d %s" % (len(before), len(after),
                        "✓✓ 完全一致" if before == after else "✗ 有变化"))

print()
print("=== 第 21 页最终文本（看长度）===")
for sh in list(prs2.slides)[20].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                print("   [%2d 字] %s" % (len(t), t[:104]))
