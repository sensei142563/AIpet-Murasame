# -*- coding: utf-8 -*-
"""按核实结果改 PPT 第 21 页措辞（先 dump 当前文本，再按实际内容改；图片位置尺寸断言不变）。"""
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

DECK = r"E:\university\order\python\pet\AIpet-Murasame-main\dev\ppt\out\Deck.pptx"
prs = Presentation(DECK)
slides = list(prs.slides)
slide = slides[20]

print("=== 第 21 页当前文本 ===")
cur = []
for sh in slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                cur.append(t)
                print("   · " + t[:118])

before = [[sh.left, sh.top, sh.width, sh.height] for s in slides for sh in s.shapes
          if sh.shape_type == 13]

# 按"实际内容"做替换：用子串匹配，容错主人可能微调过的用词
RULES = [
    ("（已落地一半）", "（代码链路已通，真机待验）"),
    ("已实现：手部 21 关键点检测", "代码已通：手部 21 关键点检测"),
    ("已实现：视线跟随", "代码已通：视线跟随"),
    ("已实现：三个手势", "代码已通：三个手势"),
    ("还没做：真机调优", "未验证：真机"),
]
NOTE = ("；相机帧 → 关键点 → 命中她 → 反应，整条链每环都有调用点"
        "（真机要用真手，须先关掉界面上的「试手」；VRM 模型还得 adb push 进去）")

n = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        t = "".join(r.text for r in para.runs)
        if not t.strip():
            continue
        newt = t
        for old, new in RULES:
            if old in newt:
                newt = newt.replace(old, new)
        # 第一条末尾补上"怎么才能真看到"
        if newt.startswith("代码已通：手部 21 关键点检测") and "试手" not in newt:
            newt = newt.rstrip() + NOTE
        if newt != t:
            if len(para.runs) == 1:
                para.runs[0].text = newt
            else:
                para.runs[0].text = newt
                for r in para.runs[1:]:
                    r.text = ""
            print("   ✓ 改为：" + newt[:104])
            n += 1
print("   共改 %d 段" % n)
if n:
    prs.save(DECK)

prs2 = Presentation(DECK)
after = [[sh.left, sh.top, sh.width, sh.height] for s in prs2.slides for sh in s.shapes
         if sh.shape_type == 13]
print()
print("=== 图片断言 ===")
print("   改前 %d / 改后 %d → %s" % (len(before), len(after),
                                    "✓✓ 完全一致" if before == after else "✗ 有变化"))

print()
print("=== 改后第 21 页 ===")
for sh in list(prs2.slides)[20].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                print("   · " + t[:118])
